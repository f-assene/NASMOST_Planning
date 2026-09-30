"""
Convertit les emplois du temps officiels (.docx ou .pdf) du dépôt en fichiers .ics lus par
l'application NASMOST Planning.

Organisation attendue du dépôt :
    <DEPARTEMENT>/<emploi du temps>.docx|.pdf     ex: NMSI/Projet d'Emp. du Temps 3GNM_15_21_Juin. 26-6.docx
Les dates de la semaine sont lues dans le nom du fichier :
    - format numérique  JJ_MM_JJ_MM_AAAA          ex: EDT_3GNM_22_06_28_06_2026.pdf
    - format officiel   JJ_JJ_Mois. AA  ou  JJ_Mois_JJ_Mois. AA
Le niveau vient du nom de la classe dans le fichier (ex: "3GNM" -> N3).

Pour chaque département, niveau et semaine, un seul fichier est produit à côté des sources :
    <DEPARTEMENT>/<DEPARTEMENT>_N<niveau>_JJ_MM_JJ_MM_AAAA.ics
Quand une même semaine existe en .docx et en .pdf, le .docx est utilisé (texte exact).

Usage : python tools/timetable_to_ics.py [racine_du_depot]
"""

import datetime as dt
import re
import sys
import unicodedata
from pathlib import Path

# Dossiers du dépôt qui ne contiennent pas d'emplois du temps.
IGNORED_DIRS = {"logos", "announcements", "annonces", "tools", ".github", ".git"}
GENERATED_MARK = "X-NASMOST-GENERATED:timetable_to_ics"
TIMEZONE = "Africa/Douala"
DEFAULT_LOCATION = "Campus Ebouyè"

DAYS = ["lundi", "mardi", "mercredi", "jeudi", "vendredi", "samedi", "dimanche"]
MONTHS = {
    "jan": 1, "janv": 1, "janvier": 1, "fev": 2, "fevr": 2, "fevrier": 2, "mar": 3, "mars": 3,
    "avr": 4, "avril": 4, "mai": 5, "juin": 6, "jun": 6, "jul": 7, "juil": 7, "juillet": 7,
    "aou": 8, "aout": 8, "sep": 9, "sept": 9, "septembre": 9, "oct": 10, "octobre": 10,
    "nov": 11, "novembre": 11, "dec": 12, "decembre": 12,
}
NUMERIC_DATES = re.compile(r"(\d{2})[-_](\d{2})[-_](\d{2})[-_](\d{2})[-_](\d{4})")
FRENCH_DATES = re.compile(
    r"(?<!\d)(\d{1,2})[\s_-]+(?:([a-z]+)\.?[\s_-]+)?(\d{1,2})[\s_-]+([a-z]+)\.?[\s_-]+(\d{4}|\d{2})(?!\d)"
)
LEVEL = re.compile(r"(?<![\dA-Za-z])([1-5])\s*(?:e|ème|eme)?\s*[A-Z]{2,5}(?![A-Za-z])")
TIME = re.compile(r"(\d{1,2})\s*[hH:]\s*(\d{2})")
CODE = re.compile(r"^[A-Z]{2,5}-[0-9X]{2,5}$")
HOURS = re.compile(r"\b(CM|TD|TP|TPE)\s*:")
TEACHER = re.compile(r"^(/|Dr\b|Dr\.|Pr\b|Pr\.|Prof|Col\b|Col\.|Cdt|Cpt|Capt|Lt|Ing|M\.|Mme|Mr\b|Mlle)")
ROOM = re.compile(r"^(salle|hall|amphi|labo|laboratoire|atelier|campus|site|bâtiment|batiment)\b", re.I)
FREE = ("ETUDES", "FERIE", "CONGE", "VACANCES")


def plain(text: str) -> str:
    """Minuscules sans accents, pour comparer les noms de jours et de mois."""
    return "".join(c for c in unicodedata.normalize("NFD", text) if unicodedata.category(c) != "Mn").lower()


def week_dates(name: str):
    """(début, fin) de la semaine d'après le nom du fichier, ou None."""
    m = NUMERIC_DATES.search(name)
    if m:
        d1, m1, d2, m2, y = map(int, m.groups())
        try:
            start = dt.date(y, m1, d1)
            end = dt.date(y + 1 if m2 < m1 else y, m2, d2)
            if end >= start:
                return start, end
        except ValueError:
            pass
    for m in FRENCH_DATES.finditer(plain(name)):
        d1, m1_text, d2, m2_text, y_text = m.groups()
        end_month = MONTHS.get(m2_text)
        if end_month is None:
            continue
        if m1_text:
            start_month = MONTHS.get(m1_text)
            if start_month is None:
                continue
        else:
            start_month = end_month if int(d1) <= int(d2) else (end_month + 10) % 12 + 1
        end_year = int(y_text) + (2000 if len(y_text) == 2 else 0)
        start_year = end_year - 1 if start_month > end_month else end_year
        try:
            start, end = dt.date(start_year, start_month, int(d1)), dt.date(end_year, end_month, int(d2))
        except ValueError:
            continue
        if end >= start:
            return start, end
    return None


def level_of(name: str) -> str:
    m = LEVEL.search(name)
    return f"N{m.group(1)}" if m else "N1"


# --- Lecture des tableaux ---------------------------------------------------------------------

def read_docx(path: Path):
    import docx  # python-docx

    document = docx.Document(str(path))
    text = "\n".join(p.text for p in document.paragraphs)
    for section in document.sections:
        text += "\n" + "\n".join(p.text for p in section.header.paragraphs)
    for table in document.tables:
        rows = [[cell.text for cell in row.cells] for row in table.rows]
        if rows and any("lundi" in plain(c) for c in rows[0]):
            return rows, text
    return None, text


def read_pdf(path: Path):
    import pdfplumber

    with pdfplumber.open(str(path)) as pdf:
        page = pdf.pages[0]
        # Tolérance fine : sans elle, les espaces entre les mots des cellules disparaissent.
        table = page.extract_table({"text_x_tolerance": 1})
        text = page.extract_text(x_tolerance=1) or ""
    return table, text


def default_room(text: str) -> str:
    """Salle indiquée dans l'en-tête ("SALLE: Salle 6 / Hall des Départements")."""
    m = re.search(r"SALLE\s*:\s*([^\n]+)", text, re.I)
    return m.group(1).strip() if m else ""


# --- Construction des événements --------------------------------------------------------------

def parse_cell(text: str):
    lines = [re.sub(r"\s+", " ", l).strip() for l in text.replace("\r", "").split("\n")]
    lines = [l for l in lines if l]
    code = lines.pop(0) if lines and CODE.match(lines[0]) else ""
    room = lines.pop() if lines and ROOM.match(lines[-1]) else ""
    teachers, hours, name = [], [], []
    for line in lines:
        if HOURS.search(line):
            hours.append(line)
        elif TEACHER.match(line) or (teachers and teachers[-1].endswith("/")):
            teachers.append(line)
        else:
            name.append(line)
    return {
        "code": code,
        "name": " ".join(name) or code or "Cours",
        "teachers": " ".join(teachers).strip(),
        "hours": " ".join(hours),
        "room": room,
    }


def escape(value: str) -> str:
    return value.replace("\\", "\\\\").replace(";", "\\;").replace(",", "\\,").replace("\n", "\\n")


def fold(line: str) -> str:
    """Coupe les lignes longues (RFC 5545 : 75 octets maximum)."""
    out, current = [], b""
    for ch in line:
        b = ch.encode("utf-8")
        if len(current) + len(b) > (75 if not out else 74):
            out.append(current.decode("utf-8"))
            current = b""
        current += b
    out.append(current.decode("utf-8"))
    return "\r\n ".join(out)


def events_for(rows, monday: dt.date, room_default: str):
    header = [plain(c) for c in rows[0]]
    columns = {}
    for index, cell in enumerate(header):
        for offset, day in enumerate(DAYS):
            if day in cell:
                columns[index] = offset
    for row in rows[1:]:
        times = TIME.findall(row[0] or "")
        if len(times) < 2:
            continue  # ligne de pause ("30 min") ou sans horaire
        (h1, m1), (h2, m2) = times[0], times[1]
        for index, offset in columns.items():
            if index >= len(row):
                continue
            cell = (row[index] or "").strip()
            if not cell or plain(cell).strip() == "pause":
                continue
            day = monday + dt.timedelta(days=offset)
            start = dt.datetime.combine(day, dt.time(int(h1), int(m1)))
            end = dt.datetime.combine(day, dt.time(int(h2), int(m2)))
            yield start, end, cell, room_default


def build_ics(dept: str, level: str, source: Path, rows, text: str, monday: dt.date) -> str:
    stamp = dt.datetime.now(dt.timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    room_default = default_room(text)
    lines = [
        "BEGIN:VCALENDAR",
        "VERSION:2.0",
        "PRODID:-//ENSTMO NASMOST Planning//timetable_to_ics//FR",
        GENERATED_MARK,
        f"X-NASMOST-SOURCE:{source.name}",
    ]
    for start, end, cell, room_fallback in events_for(rows, monday, room_default):
        uid = f"{dept}-{level}-{start:%Y%m%dT%H%M}@nasmost-planning".lower()
        lines += [
            "BEGIN:VEVENT",
            f"UID:{uid}",
            f"DTSTAMP:{stamp}",
            f"DTSTART;TZID={TIMEZONE}:{start:%Y%m%dT%H%M%S}",
            f"DTEND;TZID={TIMEZONE}:{end:%Y%m%dT%H%M%S}",
            f"CATEGORIES:{dept}-{level}",
        ]
        if any(word in plain(cell).upper() for word in FREE):
            label = "FERIE" if "FERIE" in plain(cell).upper() else "ETUDES"
            lines += [f"SUMMARY:{label}", "DESCRIPTION:Heures creuses", "END:VEVENT"]
            continue

        info = parse_cell(cell)
        description = f"Cours de {info['name']}" + (f" ({info['code']})" if info["code"] else "")
        if info["teachers"]:
            description += f"\navec {info['teachers']}"
        if info["hours"]:
            description += f"\n{info['hours']}"
        room = info["room"] or room_fallback
        lines += [
            f"SUMMARY:{escape(info['name'])}",
            f"DESCRIPTION:{escape(description)}",
            f"LOCATION:{escape(DEFAULT_LOCATION + (', ' + room if room else ''))}",
        ]
        reminders = [30, 15] if start.hour < 12 else [15]
        for minutes in reminders:
            lines += [
                "BEGIN:VALARM",
                f"TRIGGER:-PT{minutes}M",
                "ACTION:DISPLAY",
                f"DESCRIPTION:{escape(info['name'])} dans {minutes} min",
                "END:VALARM",
            ]
        lines.append("END:VEVENT")
    lines.append("END:VCALENDAR")
    return "\r\n".join(fold(l) for l in lines) + "\r\n"


# --- Parcours du dépôt ------------------------------------------------------------------------

def main(root: Path) -> int:
    written, errors = 0, 0
    for dept_dir in sorted(p for p in root.iterdir() if p.is_dir() and p.name.lower() not in IGNORED_DIRS):
        dept = dept_dir.name.upper()
        # (niveau, début de semaine) -> meilleure source ; .docx préféré au .pdf
        sources = {}
        for path in sorted(dept_dir.iterdir()):
            if path.suffix.lower() not in (".docx", ".pdf") or path.name.startswith("~$"):
                continue
            dates = week_dates(path.name)
            if not dates:
                print(f"  ignoré (dates introuvables dans le nom) : {path.relative_to(root)}")
                continue
            key = (level_of(path.name), dates)
            best = sources.get(key)
            if best is None or (path.suffix.lower() == ".docx" and best.suffix.lower() != ".docx"):
                sources[key] = path

        wanted = set()
        for (level, (start, end)), path in sorted(sources.items()):
            target = dept_dir / f"{dept}_{level}_{start:%d_%m}_{end:%d_%m_%Y}.ics"
            wanted.add(target.name)
            try:
                rows, text = read_docx(path) if path.suffix.lower() == ".docx" else read_pdf(path)
                if not rows:
                    raise ValueError("tableau de l'emploi du temps introuvable")
                monday = start - dt.timedelta(days=start.weekday())
                ics = build_ics(dept, level, path, rows, text, monday)
                # newline="" : garder les fins de ligne telles quelles pour comparer à l'identique.
                previous = open(target, encoding="utf-8", newline="").read() if target.exists() else ""
                if strip_stamps(previous) != strip_stamps(ics):
                    target.write_text(ics, encoding="utf-8", newline="")
                    written += 1
                    print(f"  écrit : {target.relative_to(root)}  (source : {path.name})")
            except Exception as error:  # un fichier illisible ne bloque pas les autres
                errors += 1
                print(f"  ERREUR {path.relative_to(root)} : {error}")

        # Supprime les .ics générés dont la source a disparu.
        for old in dept_dir.glob("*.ics"):
            if old.name not in wanted and GENERATED_MARK in old.read_text(encoding="utf-8", errors="ignore"):
                old.unlink()
                print(f"  supprimé : {old.relative_to(root)}")
    print(f"{written} fichier(s) .ics mis à jour, {errors} erreur(s).")
    return 1 if errors else 0


def strip_stamps(ics: str) -> str:
    """Contenu sans les DTSTAMP, pour ne réécrire un fichier que si les cours ont changé."""
    return re.sub(r"DTSTAMP:\d{8}T\d{6}Z", "", ics).replace("\r\n", "\n")


if __name__ == "__main__":
    sys.exit(main(Path(sys.argv[1] if len(sys.argv) > 1 else ".").resolve()))
